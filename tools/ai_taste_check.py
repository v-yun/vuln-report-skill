#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ai_taste_check — 报告去AI味机器审计（零 LLM，纯计数）
用法:
  python ai_taste_check.py <草稿.md|草稿.txt|草稿.docx>

口径全部来自 skills/report/SKILL.md「去AI腔硬规」（词面级+结构级九条）。
任一 FAIL 项超标 → 退出码 1；WARN 只提示不卡。PoC 请求块行（含 HTTP/Host:/
GET/POST 开头）不计入 ID 复现检查——那是证据原文，复现是合规的。
"""
import io, re, sys, os

# ── 阈值（与 SKILL.md 结构级九条一一对应） ──
LIMITS = {
    "connect":  ("八股连接词句首(首先/其次/综上/值得注意/也就是说…)", 0),
    "parens":   ("圆括号注解总数(≤8处且每处≤15字)", 8),
    "attacker": ("『攻击者』出现次数", 5),
    "gaix":     ("『该漏洞/该接口/该参数…』教科书腔", 6),
    "nominal":  ("『进行/实现+名词』名词化动词", 2),
    "summary":  ("『由此可见/综上』节尾总结句", 0),
    "triple":   ("『零X、零Y、零Z』三连同构排比", 0),
    "idrepeat": ("同一长ID完整复现(叙事文本内)", 0),
    "empty":    ("无数字支撑的空话套话(严重影响/极大风险…)", 0),
    "en_ism":   ("英文 model-isms(英文稿才检查)", 0),
}

CONNECT_START = ("首先", "其次", "再次", "最后，", "综上所述", "总而言之", "值得注意",
                 "需要注意", "也就是说", "换句话说", "与此同时", "进一步地", "更进一步")
NOMINAL_RE = re.compile(r"(?:进行|实现)(?:了|了一次|了一次的)?(?:验证|测试|分析|读取|操作|检查|利用|爆破|枚举|提交|访问|比对|复现)")
GAI_RE = re.compile(r"该(?:漏洞|接口|参数|系统|用户|账户|订单|功能|请求|响应|字段|平台|服务)")
SUMMARY_RE = re.compile(r"(?:由此可见|综上|可以看出)[^。！？\n]{0,30}[。！？]")
TRIPLE_RE = re.compile(r"(?:零|无|不需要?|没有)[一-龥A-Za-z]{1,6}(?:、(?:零|无|不需要?|没有)[一-龥A-Za-z]{1,6}){2,}")
EMPTY_RE = re.compile(r"(?:造成严重|带来严重|严重影响|极大风险|巨大损失|潜在威胁|不容忽视|构成严重威胁)")
EN_ISM_RE = re.compile(r"(?i)\b(?:the aforementioned|please note|furthermore|moreover|additionally|it is worth noting|in conclusion|firstly|secondly|lastly)\b")
ID_RE = re.compile(r"\b\d{7,}\b")
PARENS_RE = re.compile(r"[（(]([^（()）]*)[)）]")
POC_LINE_RE = re.compile(r"^(?:GET|POST|PUT|DELETE|PATCH|OPTIONS|HEAD)\s|^(?:Host|Content-Type|Authorization|User-Agent|Cookie|Content-Length|Accept|Connection|Set-Cookie|Server|Date|X-)\s*[a-zA-Z-]*\s*:|^HTTP/[\d.]|^https?://")

def extract_text(path: str) -> str:
    if path.lower().endswith(".docx"):
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    return open(path, encoding="utf-8", errors="replace").read()

def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(2)
    path = sys.argv[1]
    if not os.path.exists(path):
        print("文件不存在:", path); sys.exit(2)
    text = extract_text(path)

    ascii_ratio = sum(1 for c in text if ord(c) < 128) / max(1, len(text))

    # 分句（句首连接词检查用）
    sentences = [s.strip() for s in re.split(r"[。！？\n]+", text) if s.strip()]
    connect_hits = [s[:18] for s in sentences if any(s.startswith(c) for c in CONNECT_START)]

    parens = PARENS_RE.findall(text)
    # 代码引用括子(含下划线/引号/点号的技术 token)不计——那是正常技术写法不是注解
    parens = [p for p in parens if not ("_" in p or '"' in p or "`" in p or p.count(".") >= 2)]
    parens_long = [p for p in parens if len(p) > 15]
    # 英文技术文天然多用括号引代码，限额放宽到 15(仍拦得住 43 处的注解堆)
    paren_limit = 15 if ascii_ratio > 0.4 else 8

    attacker = text.count("攻击者")
    gaix = len(GAI_RE.findall(text))
    nominal = [(m.group(0), text[max(0, m.start()-10):m.end()+10]) for m in NOMINAL_RE.finditer(text)]
    summary = SUMMARY_RE.findall(text)
    triple = TRIPLE_RE.findall(text)

    # ID 复现：跳过 PoC 块行
    ids = {}
    for line in text.split("\n"):
        if POC_LINE_RE.match(line.strip()):
            continue
        for m in ID_RE.finditer(line):
            ids[m.group(0)] = ids.get(m.group(0), 0) + 1
    idrepeat = {k: v for k, v in ids.items() if v >= 2}

    # 空话套话：句内无任何数字才算违规
    empty_hits = []
    for m in EMPTY_RE.finditer(text):
        seg = text[max(0, m.start()-40):m.end()+40]
        if not re.search(r"\d", seg):
            empty_hits.append(m.group(0))

    en_ism = EN_ISM_RE.findall(text) if ascii_ratio > 0.4 else []

    # 句长均匀度（仅 WARN）
    lens = [len(s) for s in sentences if len(s) > 3]
    uniform_note = ""
    if len(lens) >= 10:
        mean = sum(lens) / len(lens)
        var = sum((l - mean) ** 2 for l in lens) / len(lens)
        sd = var ** 0.5
        if sd < 5:
            uniform_note = f"WARN 句长过均匀(mean={mean:.0f}, sd={sd:.1f}<5)——每句一样长是 AI 指纹，改出长短差"

    results = [
        ("connect", len(connect_hits), connect_hits[:3]),
        ("parens", len(parens), (f"其中超15字的: {parens_long[:3]}" if parens_long else ""), paren_limit),
        ("attacker", attacker, ""),
        ("gaix", gaix, ""),
        ("nominal", len(nominal), [n[0] for n in nominal[:3]]),
        ("summary", len(summary), summary[:2]),
        ("triple", len(triple), triple[:2]),
        ("idrepeat", len(idrepeat), dict(list(idrepeat.items())[:3])),
        ("empty", len(empty_hits), empty_hits[:3]),
        ("en_ism", len(en_ism), en_ism[:3]),
    ]
    print(f"ai_taste_check :: {os.path.basename(path)}  ({len(sentences)}句 / {len(text)}字)")
    print("-" * 72)
    fail = 0
    for row in results:
        key, actual, detail = row[0], row[1], row[2]
        name, limit = LIMITS[key]
        if len(row) > 3: limit = row[3]
        ok = actual <= limit
        if not ok: fail += 1
        mark = "PASS" if ok else "FAIL"
        print(f"[{mark}] {name:<58} {actual:>3}/{limit}" + (f"  {detail}" if detail and not ok else ""))
    if parens and len(parens) <= paren_limit and parens_long:
        print(f"[WARN] 括号总数合规但存在 >15 字长注解: {parens_long[:2]}")
    if uniform_note:
        print(uniform_note)
    print("-" * 72)
    print("结果:", "全绿" if fail == 0 else f"{fail} 项超标，回改写 pass 再测")
    sys.exit(1 if fail else 0)

if __name__ == "__main__":
    main()
